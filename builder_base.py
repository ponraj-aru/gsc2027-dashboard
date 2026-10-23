import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
# Set 16:9 widescreen dimensions
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_slide_layout = prs.slide_layouts[6]

# Theme Colors
NAVY_DARK = RGBColor(11, 19, 43)        # #0B132B
NAVY_CARD = RGBColor(22, 32, 59)        # #16203B
WHITE = RGBColor(255, 255, 255)
OFF_WHITE = RGBColor(248, 250, 252)     # #F8FAFC
CARD_BG = RGBColor(255, 255, 255)
BORDER_GRAY = RGBColor(226, 232, 240)   # #E2E8F0
TEXT_MAIN = RGBColor(15, 23, 42)        # #0F172A
TEXT_MUTED = RGBColor(100, 116, 139)    # #64748B
TEXT_LIGHT = RGBColor(148, 163, 184)   # #94A3B8
BLUE_ACCENT = RGBColor(37, 99, 235)     # #2563EB
TEAL_ACCENT = RGBColor(13, 148, 136)    # #0D9488
PURPLE_ACCENT = RGBColor(124, 58, 237)  # #7C3AED (Theme 3)
GREEN_ACCENT = RGBColor(5, 150, 105)    # #059669
AMBER_ACCENT = RGBColor(217, 119, 6)    # #D97706

def add_header(slide, title_text, category_text="GRONINGEN SUSTAINABILITY CONFERENCE 2027", is_dark=False):
    # Header box
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    # Category / Super-title
    p_cat = tf.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(9.5)
    p_cat.font.bold = True
    p_cat.font.color.rgb = BLUE_ACCENT if not is_dark else RGBColor(56, 189, 248)
    
    # Main Slide Title
    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_MAIN if not is_dark else WHITE
    p_title.space_before = Pt(3)

def set_slide_background(slide, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_GRAY):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card

print("Helper functions initialized.")
